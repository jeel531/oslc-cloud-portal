# pyrefly: ignore-errors
# type: ignore
from __future__ import annotations

import json
import os
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional
import pyodbc
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fastapi import FastAPI, Query, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, Response

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
FRONTEND_DIR = ROOT / "frontend"
CONFIG_FILE = HERE / "digi_config.json"
USERS_FILE = HERE / "users.json"

app = FastAPI(title="OSLC GODAUN API", docs_url=None, redoc_url=None)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.middleware("http")
async def add_no_cache_headers(request: Request, call_next):
    # Detect Ctrl+F5 or hard refresh from client headers or query
    client_cc = (request.headers.get("cache-control") or "").lower()
    client_prg = (request.headers.get("pragma") or "").lower()
    if "no-cache" in client_cc or "no-cache" in client_prg or request.query_params.get("refresh") in ("true", "1") or "_t" in request.query_params:
        _cache.clear()

    response = await call_next(request)
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

# In-memory caching
_cache: dict[str, tuple[float, Any]] = {}
CACHE_TTL = 30  # 30 seconds

def get_cache(key: str) -> Optional[Any]:
    if key in _cache:
        timestamp, data = _cache[key]
        if time.time() - timestamp < CACHE_TTL:
            return data
    return None

def set_cache(key: str, data: Any):
    _cache[key] = (time.time(), data)

def load_config() -> dict[str, str]:
    if not CONFIG_FILE.exists():
        raise RuntimeError("digi_config.json not found.")
    return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))

def get_db_connection():
    cfg = load_config()
    server = cfg["server"]
    database = cfg["database"]
    username = cfg["username"]
    password = cfg["password"]
    
    installed_drivers = pyodbc.drivers()
    preferred = [
        "ODBC Driver 18 for SQL Server",
        "ODBC Driver 17 for SQL Server",
        "ODBC Driver 13 for SQL Server",
        "ODBC Driver 11 for SQL Server",
        "SQL Server Native Client 11.0",
        "SQL Server Native Client 10.0",
        "SQL Server",
    ]
    driver = None
    for d in preferred:
        if d in installed_drivers:
            driver = d
            break
    if not driver:
        driver = installed_drivers[0] if installed_drivers else "SQL Server"

    params = [
        f"DRIVER={{{driver}}}",
        f"SERVER={server}",
        f"DATABASE={database}",
        f"UID={username}",
        f"PWD={password}",
        "TrustServerCertificate=yes",
    ]
    conn_str = ";".join(params)
    conn = pyodbc.connect(conn_str, timeout=15)
    conn.autocommit = True
    try:
        cur = conn.cursor()
        cur.execute("SET TRANSACTION ISOLATION LEVEL READ UNCOMMITTED;")
        cur.close()
    except Exception:
        pass
    return conn

def safe_float(val: Any) -> float:
    try:
        return round(float(val or 0), 1)
    except Exception:
        return 0.0

def load_users() -> list[dict[str, Any]]:
    if not USERS_FILE.exists():
        initial = [{
            "username": "admin",
            "password": "admin",
            "worker_name": "ADMIN",
            "role": "admin",
            "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }]
        USERS_FILE.write_text(json.dumps(initial, indent=2, ensure_ascii=False), encoding="utf-8")
        return initial
    try:
        return json.loads(USERS_FILE.read_text(encoding="utf-8"))
    except Exception:
        return []

def save_users(users: list[dict[str, Any]]):
    USERS_FILE.write_text(json.dumps(users, indent=2, ensure_ascii=False), encoding="utf-8")

@app.get("/api/godaun/users")
def get_users():
    users = load_users()
    return {"success": True, "users": users}

@app.post("/api/godaun/users")
async def save_user(req: Request):
    data = await req.json()
    old_username = str(data.get("old_username", "")).strip()
    username = str(data.get("username", "")).strip()
    password = str(data.get("password", "")).strip()
    worker_name = str(data.get("worker_name", "")).strip().upper()
    role = str(data.get("role", "worker")).strip()

    if not username or not password:
        raise HTTPException(status_code=400, detail="Username and password are required")
    if role == "worker" and not worker_name:
        raise HTTPException(status_code=400, detail="Worker name is required for worker role")

    users = load_users()

    # Check if this is an edit with a renamed username
    if old_username and old_username.lower() != username.lower():
        # Make sure target username is not already taken by someone else
        target_exists = any(u for u in users if u["username"].lower() == username.lower() and u["username"].lower() != old_username.lower())
        if target_exists:
            raise HTTPException(status_code=400, detail=f"Username '{username}' is already in use by another user")

        target_user = next((u for u in users if u["username"].lower() == old_username.lower()), None)
        if target_user:
            target_user["username"] = username
            target_user["password"] = password
            target_user["worker_name"] = worker_name
            target_user["role"] = role
            target_user["updated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            save_users(users)
            return {"success": True, "message": f"User {username} updated successfully", "users": users}

    existing = next((u for u in users if u["username"].lower() == username.lower()), None)
    if existing:
        existing["password"] = password
        existing["worker_name"] = worker_name
        existing["role"] = role
        existing["updated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    else:
        users.append({
            "username": username,
            "password": password,
            "worker_name": worker_name,
            "role": role,
            "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        })
    save_users(users)
    return {"success": True, "message": f"User {username} saved successfully", "users": users}

@app.delete("/api/godaun/users/{username}")
def delete_user(username: str):
    users = load_users()
    if username.lower() == "admin":
        raise HTTPException(status_code=400, detail="Cannot delete default admin user")
    new_users = [u for u in users if u["username"].lower() != username.lower()]
    save_users(new_users)
    return {"success": True, "message": f"User {username} deleted", "users": new_users}

@app.post("/api/godaun/login")
async def login_api(req: Request):
    data = await req.json()
    raw_u = str(data.get("username", "")).strip()
    raw_p = str(data.get("password", "")).strip()

    if not raw_u or not raw_p:
        return {"success": False, "error": "કૃપા કરીને આઈડી અને પાસવર્ડ દાખલ કરો."}

    clean_u = raw_u.lower()
    users = load_users()

    # 1. Match against existing users in users.json
    matched = next((u for u in users if (
        u.get("username", "").strip().lower() == clean_u or
        u.get("worker_name", "").strip().lower() == clean_u
    ) and str(u.get("password", "")).strip() == raw_p), None)

    if matched:
        return {"success": True, "user": matched, "message": "Login successful"}

    # Check if username exists but password was wrong
    pwd_mismatch = next((u for u in users if (
        u.get("username", "").strip().lower() == clean_u or
        u.get("worker_name", "").strip().lower() == clean_u
    )), None)
    if pwd_mismatch:
        return {"success": False, "error": "ખોટો પાસવર્ડ! કૃપા કરીને સાચો પાસવર્ડ દાખલ કરો."}

    # 2. Check if username maps to an ERP Worker in View_Dboard_Common_Dashboard
    try:
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("SELECT DISTINCT CREATED_BY FROM View_Dboard_Common_Dashboard WHERE CREATED_BY IS NOT NULL AND CREATED_BY <> ''")
        db_workers = [r[0].strip().upper() for r in cur.fetchall() if r[0] and r[0].strip()]
        conn.close()

        # Check if username matches or starts with any ERP worker
        matched_worker = next((w for w in db_workers if w.lower() == clean_u or clean_u.startswith(w.lower()) or w.lower().startswith(clean_u)), None)
        if matched_worker:
            new_user = {
                "username": raw_u.lower(),
                "password": raw_p,
                "worker_name": matched_worker,
                "role": "worker",
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            users.append(new_user)
            save_users(users)
            return {"success": True, "user": new_user, "message": "User registered and logged in successfully"}
    except Exception as db_err:
        pass

    return {"success": False, "error": "ખોટો યુઝર આઈડી અથવા પાસવર્ડ! કૃપા કરીને તપાસીને ફરી પ્રયાસ કરો."}


@app.get("/api/godaun/workers-list")
def get_workers_list():
    cached = get_cache("workers_list")
    if cached is not None:
        return cached

    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("""
        SELECT DISTINCT CREATED_BY 
        FROM View_Dboard_Common_Dashboard 
        WHERE CREATED_BY IS NOT NULL AND CREATED_BY <> ''
        ORDER BY CREATED_BY ASC
    """)
    workers = [r[0].strip() for r in cur.fetchall() if r[0] and r[0].strip()]
    conn.close()
    res = {"success": True, "workers": workers}
    set_cache("workers_list", res)
    return res

@app.get("/api/godaun/summary")
def get_summary():
    cached = get_cache("summary")
    if cached is not None:
        return cached

    today_str = datetime.now().strftime("%Y-%m-%d")
    month_start = datetime.now().strftime("%Y-%m-01")

    conn = get_db_connection()
    cur = conn.cursor()

    # Query today's and this month's stats
    query = """
        SELECT 
            TRANS_TYPE_NAME,
            SUM(CASE WHEN TRANS_DATE = ? THEN QTY_PIECES ELSE 0 END) AS TODAY_QTY,
            SUM(CASE WHEN TRANS_DATE >= ? THEN QTY_PIECES ELSE 0 END) AS MONTH_QTY
        FROM View_Dboard_Trans_All_Daily_Summary
        WHERE TRANS_DATE >= ?
          AND TRANS_TYPE_NAME IN ('WMS Stock Inward', 'WMS Order Pickup')
        GROUP BY TRANS_TYPE_NAME
    """
    cur.execute(query, (today_str, month_start, month_start))
    rows = cur.fetchall()

    today_inward = 0.0
    month_inward = 0.0
    today_outward = 0.0
    month_outward = 0.0

    for r in rows:
        t_type = r[0]
        today_val = safe_float(r[1])
        month_val = safe_float(r[2])
        if t_type == 'WMS Stock Inward':
            today_inward = today_val
            month_inward = month_val
        elif t_type == 'WMS Order Pickup':
            today_outward = today_val
            month_outward = month_val

    conn.close()

    res = {
        "success": True,
        "date": today_str,
        "today_inward": int(today_inward),
        "today_outward": int(today_outward),
        "month_inward": int(month_inward),
        "month_outward": int(month_outward),
        "last_updated": datetime.now().strftime("%H:%M:%S")
    }
    set_cache("summary", res)
    return res

@app.get("/api/godaun/reports/daywise")
def get_daywise_report(month: str = Query(default="")):
    if not month:
        month = datetime.now().strftime("%Y-%m")
    
    cache_key = f"daywise_{month}"
    cached = get_cache(cache_key)
    if cached is not None:
        return cached

    conn = get_db_connection()
    cur = conn.cursor()

    query = """
        SELECT 
            TRANS_DATE,
            TRANS_TYPE_NAME,
            SUM(QTY_PIECES) AS TOTAL_QTY
        FROM View_Dboard_Trans_All_Daily_Summary
        WHERE TRANS_DATE LIKE ? + '%'
          AND TRANS_TYPE_NAME IN ('WMS Stock Inward', 'WMS Order Pickup')
        GROUP BY TRANS_DATE, TRANS_TYPE_NAME
        ORDER BY TRANS_DATE ASC
    """
    cur.execute(query, (month,))
    rows = cur.fetchall()
    conn.close()

    days_map: dict[str, dict[str, Any]] = {}
    for r in rows:
        d_str = str(r[0])
        t_type = r[1]
        qty = safe_float(r[2])
        
        if d_str not in days_map:
            try:
                dt = datetime.strptime(d_str, "%Y-%m-%d")
                formatted = dt.strftime("%d/%m/%Y")
                short_label = dt.strftime("%d/%m")
            except Exception:
                formatted = d_str
                short_label = d_str

            days_map[d_str] = {
                "date": d_str,
                "formatted_date": formatted,
                "short_label": short_label,
                "inward": 0,
                "outward": 0
            }

        if t_type == 'WMS Stock Inward':
            days_map[d_str]["inward"] = int(qty)
        elif t_type == 'WMS Order Pickup':
            days_map[d_str]["outward"] = int(qty)

    sorted_days = sorted(days_map.values(), key=lambda x: x["date"])
    res = {
        "success": True,
        "month": month,
        "days": sorted_days,
        "total_inward": sum(d["inward"] for d in sorted_days),
        "total_outward": sum(d["outward"] for d in sorted_days),
        "last_updated": datetime.now().strftime("%H:%M:%S")
    }
    set_cache(cache_key, res)
    return res

@app.get("/api/godaun/drilldown")
def get_day_drilldown(date_str: str = Query(..., alias="date"), report_type: str = Query(default="inward")):
    cache_key = f"drill_{date_str}_{report_type}"
    cached = get_cache(cache_key)
    if cached is not None:
        return cached

    conn = get_db_connection()
    cur = conn.cursor()

    items = []
    if report_type == "inward":
        query = """
            SELECT TOP 100
                ISNULL(ITEM_NAME, 'N/A') AS ITEM_NAME,
                ISNULL(SIZE, '-') AS SIZE,
                ISNULL(SKU_CODE, 'N/A') AS SKU_CODE,
                ISNULL(ITEM_GROUP_NAME, 'GENERAL') AS CATEGORY,
                SUM(QTY_PIECES) as TOTAL_QTY
            FROM View_Dboard_Trans_All_Daily_Summary
            WHERE TRANS_DATE = ? 
              AND TRANS_TYPE_NAME = 'WMS Stock Inward'
            GROUP BY ITEM_NAME, SIZE, SKU_CODE, ITEM_GROUP_NAME
            ORDER BY TOTAL_QTY DESC
        """
        cur.execute(query, (date_str,))
        for r in cur.fetchall():
            items.append({
                "design": r[0],
                "size": r[1],
                "sku": r[2],
                "category": r[3],
                "qty": int(safe_float(r[4]))
            })
    else:
        query = """
            SELECT TOP 100
                ISNULL(BILL_NO, 'BOX') AS BOX_NO,
                ISNULL(VOUCHER_NO, '') AS VOUCHER_NO,
                SUM(QTY_PIECES) AS TOTAL_QTY,
                COUNT(*) AS ENTRIES,
                ISNULL(CREATED_BY, '') AS CREATED_BY
            FROM View_Dboard_Common_Dashboard
            WHERE TRANS_DATE = ? 
              AND PROCESS_NAME = 'Portal Box Creation (ASIN)'
            GROUP BY BILL_NO, VOUCHER_NO, CREATED_BY
            ORDER BY TOTAL_QTY DESC
        """
        cur.execute(query, (date_str,))
        for r in cur.fetchall():
            items.append({
                "box_no": r[0],
                "voucher_no": r[1],
                "qty": int(safe_float(r[2])),
                "entries": int(r[3]),
                "created_by": r[4]
            })
    conn.close()
    res = {
        "success": True,
        "date": date_str,
        "type": report_type,
        "items": items
    }
    set_cache(cache_key, res)
    return res

@app.get("/api/godaun/user-summary")
def get_user_summary(date_str: str = Query(..., alias="date"), report_type: str = Query(default="inward")):
    cache_key = f"user_summary_{date_str}_{report_type}"
    cached = get_cache(cache_key)
    if cached is not None:
        return cached

    conn = get_db_connection()
    cur = conn.cursor()

    p_name = "WMS Stock Inward" if report_type == "inward" else "WMS Order Pickup"

    query = """
        SELECT 
            ISNULL(CREATED_BY, 'UNKNOWN') AS USER_NAME,
            SUM(QTY_PIECES) AS TOTAL_QTY,
            COUNT(*) AS ENTRIES
        FROM View_Dboard_Common_Dashboard
        WHERE TRANS_DATE = ? AND PROCESS_NAME = ?
          AND (BARCODE IS NOT NULL AND BARCODE <> '')
        GROUP BY CREATED_BY
        ORDER BY TOTAL_QTY DESC
    """
    cur.execute(query, (date_str, p_name))
    rows = cur.fetchall()
    conn.close()

    total_sum = sum(safe_float(r[1]) for r in rows)
    users = []
    for r in rows:
        qty = int(safe_float(r[1]))
        pct = round((qty / total_sum * 100), 1) if total_sum > 0 else 0.0
        users.append({
            "user_name": r[0],
            "total_qty": qty,
            "entries": int(r[2]),
            "percentage": pct
        })

    res = {
        "success": True,
        "date": date_str,
        "type": report_type,
        "process_name": p_name,
        "total_qty": int(total_sum),
        "user_count": len(users),
        "users": users,
        "last_updated": datetime.now().strftime("%H:%M:%S")
    }
    set_cache(cache_key, res)
    return res

@app.get("/api/godaun/user-detail")
def get_user_detail(
    date_str: str = Query(..., alias="date"),
    report_type: str = Query(default="inward"),
    user_name: str = Query(..., alias="user")
):
    cache_key = f"user_detail_{date_str}_{report_type}_{user_name}"
    cached = get_cache(cache_key)
    if cached is not None:
        return cached

    conn = get_db_connection()
    cur = conn.cursor()

    p_name = "WMS Stock Inward" if report_type == "inward" else "WMS Order Pickup"

    query = """
        SELECT TOP 200
            ISNULL(VOUCHER_NO, '') AS VOUCHER_NO,
            ISNULL(BILL_NO, '-') AS BILL_NO,
            ISNULL(DESIGN_NO, 'N/A') AS DESIGN_NO,
            ISNULL(SIZE, '-') AS SIZE,
            QTY_PIECES,
            CREATED_DATE
        FROM View_Dboard_Common_Dashboard
        WHERE TRANS_DATE = ? 
          AND PROCESS_NAME = ?
          AND CREATED_BY = ?
          AND (BARCODE IS NOT NULL AND BARCODE <> '')
        ORDER BY CREATED_DATE DESC
    """
    cur.execute(query, (date_str, p_name, user_name))
    items = []
    for r in cur.fetchall():
        dt_str = ""
        if r[5]:
            try:
                dt_str = r[5].strftime("%H:%M:%S")
            except Exception:
                dt_str = str(r[5])
        items.append({
            "voucher_no": r[0],
            "bill_no": r[1],
            "design_no": r[2],
            "size": r[3],
            "qty": int(safe_float(r[4])),
            "time": dt_str
        })
    conn.close()

    res = {
        "success": True,
        "date": date_str,
        "type": report_type,
        "user_name": user_name,
        "count": len(items),
        "items": items
    }
    set_cache(cache_key, res)
    return res

def fetch_salary_report_data(from_date: str, to_date: str, in_rate: float, out_rate: float, name_filter: str = ""):
    conn = get_db_connection()
    cur = conn.cursor()

    query = """
        SELECT 
            ISNULL(CREATED_BY, 'UNKNOWN') AS NAME,
            SUM(CASE WHEN PROCESS_NAME = 'WMS Stock Inward' THEN QTY_PIECES ELSE 0 END) AS IN_PCS,
            SUM(CASE WHEN PROCESS_NAME = 'WMS Order Pickup' THEN QTY_PIECES ELSE 0 END) AS OUT_PCS
        FROM View_Dboard_Common_Dashboard
        WHERE TRANS_DATE >= ? AND TRANS_DATE <= ?
          AND PROCESS_NAME IN ('WMS Stock Inward', 'WMS Order Pickup')
          AND (BARCODE IS NOT NULL AND BARCODE <> '')
        GROUP BY CREATED_BY
        ORDER BY (SUM(CASE WHEN PROCESS_NAME = 'WMS Stock Inward' THEN QTY_PIECES ELSE 0 END) + SUM(CASE WHEN PROCESS_NAME = 'WMS Order Pickup' THEN QTY_PIECES ELSE 0 END)) DESC
    """
    cur.execute(query, (from_date, to_date))
    rows = cur.fetchall()
    conn.close()

    result_rows = []
    tot_in_pcs = 0
    tot_out_pcs = 0
    tot_in_amt = 0.0
    tot_out_amt = 0.0

    name_filter_lower = name_filter.strip().lower()

    for r in rows:
        name = str(r[0]).strip()
        if name_filter_lower and name_filter_lower not in name.lower():
            continue
        in_pcs = int(safe_float(r[1]))
        out_pcs = int(safe_float(r[2]))
        in_amt = round(in_pcs * in_rate, 2)
        out_amt = round(out_pcs * out_rate, 2)
        total_pc = in_pcs + out_pcs
        total_amt = round(in_amt + out_amt, 2)

        tot_in_pcs += in_pcs
        tot_out_pcs += out_pcs
        tot_in_amt += in_amt
        tot_out_amt += out_amt

        result_rows.append({
            "name": name,
            "in_pcs": in_pcs,
            "in_amt": in_amt,
            "out_pcs": out_pcs,
            "out_amt": out_amt,
            "total_pc": total_pc,
            "total_amt": total_amt
        })

    # Sort strictly descending by total_amt (highest work/earnings first, lowest last)
    result_rows.sort(key=lambda r: (r['total_amt'], r['total_pc']), reverse=True)

    totals = {
        "in_pcs": tot_in_pcs,
        "in_amt": round(tot_in_amt, 2),
        "out_pcs": tot_out_pcs,
        "out_amt": round(tot_out_amt, 2),
        "total_pc": tot_in_pcs + tot_out_pcs,
        "total_amt": round(tot_in_amt + tot_out_amt, 2)
    }

    return result_rows, totals

@app.get("/api/godaun/salary-report")
def get_salary_report(
    from_date: str = "",
    to_date: str = "",
    in_rate: float = 0.20,
    out_rate: float = 0.35,
    name: str = ""
):
    if not from_date:
        from_date = datetime.now().strftime("%Y-%m-%d")
    if not to_date:
        to_date = from_date

    cache_key = f"salary_{from_date}_{to_date}_{in_rate}_{out_rate}_{name}"
    cached = get_cache(cache_key)
    if cached is not None:
        return cached

    rows, totals = fetch_salary_report_data(from_date, to_date, in_rate, out_rate, name)
    res = {
        "success": True,
        "from_date": from_date,
        "to_date": to_date,
        "in_rate": in_rate,
        "out_rate": out_rate,
        "name_filter": name,
        "rows": rows,
        "top5": rows[:5],
        "totals": totals,
        "count": len(rows),
        "last_updated": datetime.now().strftime("%H:%M:%S")
    }
    set_cache(cache_key, res)
    return res

@app.get("/api/godaun/user-datewise")
def get_user_datewise(
    user: str = "",
    from_date: str = "",
    to_date: str = "",
    in_rate: float = 0.20,
    out_rate: float = 0.35
):
    if not from_date:
        from_date = datetime.now().strftime("%Y-%m-01")
    if not to_date:
        to_date = datetime.now().strftime("%Y-%m-%d")

    cache_key = f"datewise_{user}_{from_date}_{to_date}_{in_rate}_{out_rate}"
    cached = get_cache(cache_key)
    if cached is not None:
        return cached

    conn = get_db_connection()
    cur = conn.cursor()

    query = """
        SELECT 
            TRANS_DATE,
            SUM(CASE WHEN PROCESS_NAME = 'WMS Stock Inward' THEN QTY_PIECES ELSE 0 END) AS IN_PCS,
            SUM(CASE WHEN PROCESS_NAME = 'WMS Order Pickup' THEN QTY_PIECES ELSE 0 END) AS OUT_PCS
        FROM View_Dboard_Common_Dashboard
        WHERE TRANS_DATE >= ? AND TRANS_DATE <= ?
          AND CREATED_BY = ?
          AND PROCESS_NAME IN ('WMS Stock Inward', 'WMS Order Pickup')
          AND (BARCODE IS NOT NULL AND BARCODE <> '')
        GROUP BY TRANS_DATE
        ORDER BY TRANS_DATE ASC
    """
    cur.execute(query, (from_date, to_date, user))
    rows = cur.fetchall()
    conn.close()

    daily_rows = []
    tot_in = 0
    tot_out = 0
    tot_in_amt = 0.0
    tot_out_amt = 0.0

    for r in rows:
        d_str = str(r[0])
        try:
            formatted = datetime.strptime(d_str, "%Y-%m-%d").strftime("%d/%m/%Y")
        except Exception:
            formatted = d_str
        in_pcs = int(safe_float(r[1]))
        out_pcs = int(safe_float(r[2]))
        in_amt = round(in_pcs * in_rate, 2)
        out_amt = round(out_pcs * out_rate, 2)
        tot_in += in_pcs
        tot_out += out_pcs
        tot_in_amt += in_amt
        tot_out_amt += out_amt

        daily_rows.append({
            "date": d_str,
            "formatted_date": formatted,
            "in_pcs": in_pcs,
            "in_amt": in_amt,
            "out_pcs": out_pcs,
            "out_amt": out_amt,
            "total_pc": in_pcs + out_pcs,
            "total_amt": round(in_amt + out_amt, 2)
        })

    totals = {
        "in_pcs": tot_in,
        "in_amt": round(tot_in_amt, 2),
        "out_pcs": tot_out,
        "out_amt": round(tot_out_amt, 2),
        "total_pc": tot_in + tot_out,
        "total_amt": round(tot_in_amt + tot_out_amt, 2)
    }

    res = {
        "success": True,
        "user_name": user,
        "from_date": from_date,
        "to_date": to_date,
        "in_rate": in_rate,
        "out_rate": out_rate,
        "days": daily_rows,
        "totals": totals
    }
    set_cache(cache_key, res)
    return res

@app.get("/api/godaun/export-excel")
def export_excel_report(
    from_date: str = "",
    to_date: str = "",
    in_rate: float = 0.20,
    out_rate: float = 0.35,
    name: str = "",
    user: str = ""
):
    if not from_date:
        from_date = datetime.now().strftime("%Y-%m-%d")
    if not to_date:
        to_date = from_date

    wb = openpyxl.Workbook()
    ws = wb.active

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2C4A6E", end_color="2C4A6E", fill_type="solid")

    amt_in_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    amt_out_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    tot_amt_fill = PatternFill(start_color="F9CB9C", end_color="F9CB9C", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )

    if user:
        ws.title = f"{user[:25]} Daily"
        res_datewise = get_user_datewise(user, from_date, to_date, in_rate, out_rate)
        days = res_datewise["days"]
        totals = res_datewise["totals"]

        headers = ["DATE", "IN", f"{in_rate:.2f}", "OUT", f"{out_rate:.2f}", "TOTAL PC", "TOTAL AMT"]
        ws.append(headers)

        for col_idx in range(1, 8):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

        for r_idx, r in enumerate(days, start=2):
            ws.append([
                r["formatted_date"],
                r["in_pcs"],
                r["in_amt"],
                r["out_pcs"],
                r["out_amt"],
                r["total_pc"],
                r["total_amt"]
            ])
            for col_idx in range(1, 8):
                c = ws.cell(row=r_idx, column=col_idx)
                c.border = thin_border
                if col_idx == 3:
                    c.fill = amt_in_fill
                elif col_idx == 5:
                    c.fill = amt_out_fill
                elif col_idx == 7:
                    c.fill = tot_amt_fill

        tot_row_idx = len(days) + 2
        ws.append([
            "TOTAL",
            totals["in_pcs"],
            totals["in_amt"],
            totals["out_pcs"],
            totals["out_amt"],
            totals["total_pc"],
            totals["total_amt"]
        ])
        for col_idx in range(1, 8):
            c = ws.cell(row=tot_row_idx, column=col_idx)
            c.font = Font(name="Calibri", size=11, bold=True)
            c.border = thin_border
            if col_idx == 7:
                c.fill = PatternFill(start_color="F6B26B", end_color="F6B26B", fill_type="solid")

        filename = f"OSLC_{user}_Stock_Report_{from_date}_to_{to_date}.xlsx"
    else:
        ws.title = "Stock Report"
        rows, totals = fetch_salary_report_data(from_date, to_date, in_rate, out_rate, name)

        headers = ["NAME", "IN", f"{in_rate:.2f}", "OUT", f"{out_rate:.2f}", "TOTAL PC", "TOTAL AMT"]
        ws.append(headers)

        for col_idx in range(1, 8):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

        for r_idx, r in enumerate(rows, start=2):
            ws.append([
                r["name"],
                r["in_pcs"],
                r["in_amt"],
                r["out_pcs"],
                r["out_amt"],
                r["total_pc"],
                r["total_amt"]
            ])
            for col_idx in range(1, 8):
                c = ws.cell(row=r_idx, column=col_idx)
                c.border = thin_border
                if col_idx == 3:
                    c.fill = amt_in_fill
                elif col_idx == 5:
                    c.fill = amt_out_fill
                elif col_idx == 7:
                    c.fill = tot_amt_fill

        tot_row_idx = len(rows) + 2
        ws.append([
            "TOTAL",
            totals["in_pcs"],
            totals["in_amt"],
            totals["out_pcs"],
            totals["out_amt"],
            totals["total_pc"],
            totals["total_amt"]
        ])
        for col_idx in range(1, 8):
            c = ws.cell(row=tot_row_idx, column=col_idx)
            c.font = Font(name="Calibri", size=11, bold=True)
            c.border = thin_border
            if col_idx == 7:
                c.fill = PatternFill(start_color="F6B26B", end_color="F6B26B", fill_type="solid")

        filename = f"OSLC_Stock_Report_{from_date}_to_{to_date}.xlsx"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    out = io.BytesIO()
    wb.save(out)
    excel_bytes = out.getvalue()

    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# Mount frontend static files and index.html at root
if FRONTEND_DIR.exists():
    app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8095)
